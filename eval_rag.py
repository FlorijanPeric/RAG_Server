#!/usr/bin/env python3
"""Measure how well the RAG system answers the bee-keeping quiz questions.

Uses the Slovenian multiple-choice questions from
dataset/questions/Slovenian-English questions database_odgovori.docx, where the
correct option is coloured green. Runs against the current configuration (same
env vars as the server), so you can compare embedders / LLMs / collections.

    python eval_rag.py                   # retrieval: is the correct answer in the top-k chunks?
    python eval_rag.py --llm 40          # + let the LLM answer 40 questions (a/b/c...), with RAG context
    python eval_rag.py --llm 40 --no-context   # same questions without context (baseline)

The questions file must NOT be in the collection you evaluate (it contains the
questions and options verbatim), so ingest only dataset/data, e.g.:

    QDRANT_COLLECTION=eval python init_rag.py --reset --dir dataset/data
    QDRANT_COLLECTION=eval python eval_rag.py --llm 40
"""
import argparse
import random
import re
import sys
from pathlib import Path

from rag_pipeline import RAGPipeline

QUESTIONS_FILE = Path("dataset/questions/Slovenian-English questions database_odgovori.docx")
GREEN = "00B050"


def load_questions(path: Path) -> list[dict]:
    import docx

    document = docx.Document(str(path))
    questions = []
    for row in document.tables[0].rows:
        paragraphs = [p for p in row.cells[0].paragraphs if p.text.strip()]
        if len(paragraphs) < 3:
            continue
        options = paragraphs[1:]
        correct = [
            index
            for index, paragraph in enumerate(options)
            if any(
                run.font.color is not None
                and run.font.color.type is not None
                and str(run.font.color.rgb) == GREEN
                for run in paragraph.runs
            )
        ]
        # Only the Slovenian blocks have the answer marked (exactly one green option).
        if len(correct) != 1:
            continue
        questions.append(
            {
                "question": paragraphs[0].text.strip(),
                "options": [p.text.strip() for p in options],
                "answer": correct[0],
            }
        )
    return questions


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text.lower().replace(" ", " ")).strip(" .,;:")


def eval_retrieval(pipeline: RAGPipeline, questions: list[dict], top_k: int) -> None:
    hits = evaluated = leaked = reachable = 0
    # Upper bound: questions whose answer text appears in some chunk of the collection at all.
    index_points = pipeline._get_bm25()[1]
    all_chunks = [normalize((point.payload or {}).get("text", "")) for point in index_points]
    for item in questions:
        answer = normalize(item["options"][item["answer"]])
        # Very short answers ("1", "2") match almost any text; skip them.
        if len(answer) < 4 or answer in normalize(item["question"]):
            continue
        contexts = pipeline.retrieve(item["question"], top_k=top_k)
        if any("questions" in str(c["metadata"].get("source")) for c in contexts):
            leaked += 1
        joined = normalize(" ".join(c["text"] for c in contexts))
        evaluated += 1
        hits += answer in joined
        reachable += any(answer in chunk for chunk in all_chunks)

    print(f"retrieval: correct answer text in top-{top_k} chunks for {hits}/{evaluated} "
          f"questions ({hits / max(evaluated, 1):.0%}; the answer appears in the collection at all "
          f"for {reachable}/{evaluated} = upper bound)")
    if leaked:
        print(f"  WARNING: {leaked} queries retrieved chunks from dataset/questions - the questions "
              "file is in this collection, so these numbers are inflated. Ingest only dataset/data.")


def eval_llm(pipeline: RAGPipeline, questions: list[dict], top_k: int, use_context: bool) -> None:
    letters = "abcdef"
    correct = answered = 0
    for number, item in enumerate(questions, start=1):
        option_letters = letters[: len(item["options"])]
        options = "\n".join(f"{letters[i]}) {option}" for i, option in enumerate(item["options"]))
        if use_context:
            contexts = pipeline.retrieve(item["question"], top_k=top_k)
            context_block = "\n\n".join(f"[{i + 1}] {c['text']}" for i, c in enumerate(contexts))
            prompt = (
                "Na podlagi spodnjega konteksta odgovori na vprašanje z izbiro enega odgovora.\n\n"
                f"Kontekst:\n{context_block}\n\n"
            )
        else:
            prompt = "Odgovori na vprašanje z izbiro enega odgovora.\n\n"
        prompt += (
            f"Vprašanje: {item['question']}\n{options}\n\n"
            f"Odgovori samo s črko ({', '.join(option_letters)})."
        )
        try:
            reply = pipeline.generate(prompt)
        except RuntimeError as error:
            print(f"  [{number}] LLM error: {error}")
            continue
        match = re.search(rf"\b([{option_letters}])\b", reply.lower())
        answered += 1
        is_correct = bool(match) and letters.index(match.group(1)) == item["answer"]
        correct += is_correct
        print(f"  [{number:>3}] {'OK ' if is_correct else 'BAD'} {item['question'][:70]}")

    label = "with RAG context" if use_context else "without context"
    print(f"llm ({label}): {correct}/{answered} correct ({correct / max(answered, 1):.0%}, "
          "random guessing ~ 33%)")


def main() -> int:
    parser = argparse.ArgumentParser(description="Evaluate the RAG system on the quiz questions")
    parser.add_argument("--top-k", type=int, default=None, help="Default: TOP_K from .env")
    parser.add_argument("--llm", type=int, default=0, metavar="N",
                        help="Also let the LLM answer N randomly chosen questions")
    parser.add_argument("--no-context", action="store_true",
                        help="With --llm: ask without retrieved context (baseline)")
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()

    questions = load_questions(QUESTIONS_FILE)
    pipeline = RAGPipeline()
    print(f"{len(questions)} questions | embedder={pipeline.describe_embedder()} | "
          f"llm={pipeline.describe_llm()} | collection={pipeline.config.qdrant_collection}")
    print(f"retrieval={pipeline.describe_retrieval()}")
    if not pipeline.has_data():
        print("collection is empty - ingest first (python init_rag.py --dir dataset/data)")
        return 1

    args.top_k = args.top_k or pipeline.config.top_k
    eval_retrieval(pipeline, questions, args.top_k)
    if args.llm:
        sample = random.Random(args.seed).sample(questions, min(args.llm, len(questions)))
        eval_llm(pipeline, sample, args.top_k, use_context=not args.no_context)
    return 0


if __name__ == "__main__":
    sys.exit(main())
