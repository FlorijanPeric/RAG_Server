"""Shared file discovery/extraction for ingesting a whole dataset directory.

Supports .pdf, .txt, .docx, .xls. Used both by app.py's auto-ingest-on-startup
and as a standalone CLI for manually (re)ingesting everything:

    python ingest_dataset.py --dir dataset
"""
import argparse
from pathlib import Path

SUPPORTED_EXTENSIONS = (".pdf", ".txt", ".docx", ".xls")


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()

    if suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts)

    if suffix == ".xls":
        import pandas as pd

        sheets = pd.read_excel(path, sheet_name=None)
        parts = [f"# Sheet: {name}\n{df.to_csv(index=False)}" for name, df in sheets.items()]
        return "\n\n".join(parts)

    return path.read_text(encoding="utf-8", errors="ignore")


def discover_files(root: Path) -> list[Path]:
    files: list[Path] = []
    for ext in SUPPORTED_EXTENSIONS:
        files.extend(sorted(root.rglob(f"*{ext}")))
    return files


def ingest_directory(pipeline, root: Path) -> dict:
    files = discover_files(root)
    ingested = 0
    skipped = 0

    for path in files:
        try:
            text = extract_text(path)
            if not text.strip():
                print(f"skip (no extractable text): {path}")
                skipped += 1
                continue
            result = pipeline.ingest_text(text=text, source=str(path))
            ingested += 1
            print(f"ingested {path} -> {result.get('chunks_added', 0)} chunks")
        except Exception as error:
            print(f"failed on {path}: {error}")
            skipped += 1

    return {"files_found": len(files), "ingested": ingested, "skipped": skipped}


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Ingest all supported files (.pdf, .txt, .docx, .xls) from a directory"
    )
    parser.add_argument("--dir", "-d", default="dataset", help="Directory to scan recursively")
    args = parser.parse_args()

    root = Path(args.dir)
    if not root.exists():
        raise FileNotFoundError(f"Directory not found: {root}")

    from rag_pipeline import RAGPipeline

    pipeline = RAGPipeline()
    summary = ingest_directory(pipeline, root)
    print(summary)


if __name__ == "__main__":
    main()
